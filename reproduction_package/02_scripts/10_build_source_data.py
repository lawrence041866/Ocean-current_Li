# Build Source Data workbook (numeric values behind every figure) -> Source_Data.xlsx
import pandas as pd, numpy as np, networkx as nx, community as cl
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
D="01_data/"; OUT="03_figures/"
HDR=Font(name='Arial',bold=True,color='FFFFFF'); HFILL=PatternFill('solid',start_color='1F4E79')
BASE=Font(name='Arial'); TITLE=Font(name='Arial',bold=True,size=12)

wb=Workbook(); ws=wb.active; ws.title='Overview'
ov=[['Source Data — Ocean-current connectivity drives the transboundary synchrony of marine pollution'],
    [''],
    ['Sheet','Figure','Content'],
    ['Fig1a','Figure 1a','Spatial spillover coefficients (ocean vs geography) across four outcomes, with 95% CI'],
    ['Fig1b','Figure 1b','OLS vs spatial-2SLS connectivity coefficient, with 95% CI'],
    ['Fig2a','Figure 2a','Spatial-channel horse-race coefficients (ocean & geography), with 95% CI'],
    ['Fig2b','Figure 2b','Thermal/upwelling/first-difference robustness of the ocean coefficient, with 95% CI'],
    ['Fig3_clusters','Figure 3','Per-nation cluster assignment, name, display number and centroid coordinates'],
    ['Fig3_synchrony','Figure 3','Within- vs between-cluster chlorophyll-anomaly correlations (7.4x, p)'],
    ['']]
for r in ov: ws.append(r)
ws['A1'].font=TITLE
for c in ws[3]: c.font=HDR; c.fill=HFILL

def sheet_from_df(name,df,note=None):
    s=wb.create_sheet(name); 
    if note: s.append([note]); s['A1'].font=Font(name='Arial',italic=True); s.append([])
    s.append(list(df.columns))
    for c in s[s.max_row]: c.font=HDR; c.fill=HFILL
    for _,row in df.iterrows(): s.append(list(row.values))
    for col in s.columns:
        w=max(len(str(c.value)) for c in col if c.value is not None)+2
        s.column_dimensions[col[0].column_letter].width=min(max(w,10),60)
    for row in s.iter_rows(min_row=1):
        for c in row:
            if c.font is BASE or c.font.color is None or not c.font.bold: 
                if c.font.color is None: c.font=BASE
    return s

# Fig1
f1=pd.read_csv(OUT+"fig1_coef_ci.csv")
f1a=f1[f1.outcome.isin(['Chlorophyll-a (SDG 14.1.1)','OHI Clean Waters','OHI Index (overall)','OHI Biodiversity'])].copy()
sheet_from_df('Fig1a',f1a,'Figure 1a — panel spatial-lag models, two-way FE, country-clustered SE.')
f1b=pd.DataFrame([
 {'estimator':'OLS (ocean-only)','coef':float(f1[f1.outcome=='OLS (ocean-only, chl)'].coef.iloc[0]),
  'ci_lo':float(f1[f1.outcome=='OLS (ocean-only, chl)'].ci_lo.iloc[0]),'ci_hi':float(f1[f1.outcome=='OLS (ocean-only, chl)'].ci_hi.iloc[0]),'n':2337},
 {'estimator':'Spatial 2SLS (Kelejian-Prucha)','coef':0.98,'ci_lo':0.63,'ci_hi':1.32,'n':2337}])
sheet_from_df('Fig1b',f1b,'Figure 1b — first-stage F = 96.2; instruments = W Y(t-1), W^2 Y(t-1).')

# Fig2
f2=pd.read_csv(OUT+"fig2_coef_ci.csv")
sheet_from_df('Fig2a',f2[f2.panel=='2a'].drop(columns='panel'),'Figure 2a — horse-race; n = 2,337 nation-years.')
sheet_from_df('Fig2b',f2[f2.panel=='2b'].drop(columns='panel'),'Figure 2b — thermal/upwelling/first-difference; n = 2,074.')

# Fig3 clusters
clu=pd.read_csv(D+"c1_clusters.csv")
hx=pd.read_csv(D+"hex2country.csv")
cen=hx.groupby('iso3')[['lon','lat']].median()
names={0:'W Indian Ocean / Arabian Sea',1:'Mediterranean',2:'NW Africa',3:'SW/SE Atlantic & SW Pacific',
       4:'Wider Caribbean',5:'SE Asia / W Pacific',6:'NE Atlantic / North & Baltic',7:'W Africa / Gulf of Guinea',
       8:'N Indian Ocean / Bay of Bengal',9:'NW Pacific (+ Black Sea)',10:'Gulf of Mexico / Central America'}
order=[1,6,0,8,5,4,10,9,7,3,2]; disp={c:i+1 for i,c in enumerate(order)}
cl3=clu.copy()
cl3['cluster_name']=cl3.cluster.map(names); cl3['display_no']=cl3.cluster.map(disp)
cl3['lon']=cl3.iso3.map(cen.lon).round(2); cl3['lat']=cl3.iso3.map(cen.lat).round(2)
cl3=cl3.sort_values(['display_no','iso3'])[['iso3','cluster','cluster_name','display_no','lon','lat']]
sheet_from_df('Fig3_clusters',cl3,'Figure 3 — Louvain cluster id (algorithmic), name, legend display number, centroid.')

# Fig3 synchrony within/between
W=pd.read_csv(D+"W_sovereign_directed.csv")
chl=pd.read_csv(D+"chl_yearly_2003_2021.csv"); chl['year']=chl.year.astype(int)
chl=chl.groupby(['iso3','year']).apply(lambda g:np.average(g.chlor_a_mean,weights=g.pixel_count)).reset_index(name='chl')
chl['lchl']=np.log(chl.chl)
countries=sorted(set(chl.iso3)&(set(W.src)|set(W.dst)))
part=clu.set_index('iso3').cluster.to_dict()
piv=chl[chl.iso3.isin(countries)].pivot(index='year',columns='iso3',values='lchl')
anom=piv.apply(lambda c: c-np.polyval(np.polyfit(piv.index[c.notna()],c[c.notna()],1),piv.index) if c.notna().sum()>3 else c)
pin,pout=[],[]
cs=list(countries)
for i in range(len(cs)):
    for j in range(i+1,len(cs)):
        a,b=cs[i],cs[j]
        if a in anom and b in anom and a in part and b in part:
            dd=pd.DataFrame({'a':anom[a],'b':anom[b]}).dropna()
            if len(dd)>=10:
                r=np.corrcoef(dd.a,dd.b)[0,1]
                (pin if part[a]==part[b] else pout).append(r)
p=stats.mannwhitneyu(pin,pout,alternative='greater')[1]
syn=pd.DataFrame([
 {'group':'within-cluster pairs','n_pairs':len(pin),'mean_corr':round(np.mean(pin),4),'median_corr':round(np.median(pin),4)},
 {'group':'between-cluster pairs','n_pairs':len(pout),'mean_corr':round(np.mean(pout),4),'median_corr':round(np.median(pout),4)},
 {'group':'ratio (within/between, mean)','n_pairs':'','mean_corr':round(np.mean(pin)/np.mean(pout),2),'median_corr':''},
 {'group':'Mann-Whitney U p-value','n_pairs':'','mean_corr':f'{p:.2e}','median_corr':''}])
sheet_from_df('Fig3_synchrony',syn,'Figure 3 — within vs between cluster chlorophyll-anomaly correlations.')

wb.save("Source_Data.xlsx")
print("Source_Data.xlsx saved with sheets:", wb.sheetnames)
